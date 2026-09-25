"""Build robotai_BOM.xlsx from the verified research data below.

Every price/URL was read off the live product page (or search results page where
noted) on 2026-09-25. Re-run after editing the rows: python tools/build_bom.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink

OUT = Path(__file__).resolve().parent.parent / "robotai_BOM.xlsx"
VERIFIED = "2026-09-25"

# (Item, Product Name, Link, Unit Price, Qty, Notes)
CORE = [
    ("ESP32 dev board (x2, one per hand)",
     "Freenove ESP32 Development Board ESP32 Kit (2 Pack)",
     "https://www.amazon.ca/dp/B0C9THDPXP", 27.95, 1,
     "One 2-pack covers both hands ($13.98/board). ESP32-WROOM, 4.4* (527), well documented."),
    ("Finger servos (10 + 2 spares) - GOOD tier",
     "Deegoo-FPV MG996R Metal Gear Digital Servo, 4-Pack",
     "https://www.amazon.ca/dp/B07MFK266B", 35.22, 3,
     "12 metal-gear standard servos, ~12 kg-cm @6V, $8.81/ea, 4.5* (1,182). Fits stock InMoov forearm. "
     "Showed 'only 6 left' - fallback: Aideepen MG996R 6-pack x2 (see Alternatives)."),
    ("Servo driver (x2, one per hand)",
     "HiLetgo PCA9685 16-Channel 12-Bit PWM Servo Driver (I2C), 2-Pack",
     "https://www.amazon.ca/dp/B07BRS249H", 24.14, 1,
     "Adafruit-layout clone, works with Adafruit PWM library on ESP32. One board per hand matches 1-ESP32-per-hand; "
     "a single board (16 ch) could drive all 10 servos but would tie both hands to one ESP32."),
    ("Servo power supply (x2, one per hand)",
     "SHNITPWR 3.5-12V Adjustable 10A 120W Power Supply w/ LCD",
     "https://www.amazon.ca/dp/B08N4FB7MF", 36.99, 2,
     "Set to 6V. 10A = 5 MG996R running + 2-3 stalled (~2.5A stall each). LCD guards against accidental 12V. "
     "4.3* (216). Servos NEVER on ESP32 5V pin; tie grounds."),
    ("DC barrel-jack to screw-terminal adapters",
     "20 PCS Female DC 5.5x2.1mm Barrel to Screw Terminal",
     "https://www.amazon.ca/dp/B0C4JJ2HB4", 11.99, 1,
     "Connects PSU to PCA9685 V+ terminal. Small jacks run warm near 10A - cutting the plug and screwing wire in is safer."),
    ("Tendon line (shared spool)",
     "HERCULES 8-Strand Braided PE Fishing Line, 50lb, 0.37mm, 100m",
     "https://www.amazon.ca/dp/B0791D5G7J", 16.99, 1,
     "InMoov-class braid, low stretch; 100m is far more than both hands need. 4.5* (7.7K). Price for 50lb/100m variant."),
    ("Finger return springs",
     "Swpeet 150pc Stainless Extension Spring Assortment (15 sizes x 10)",
     "https://www.amazon.ca/dp/B0D7W3RT31", 20.99, 1,
     "10 of each size, so 10 matching springs per size for 10 fingers, lets you tune return force. 4.6* (111)."),
    ("Jumper wire kit",
     "ELEGOO 120pc 20cm Dupont Wire (40 M-F / 40 M-M / 40 F-F)",
     "https://www.amazon.ca/dp/B01EV70C78", 12.99, 1,
     "Signal/I2C wiring ESP32<->PCA9685. 4.7* (14.8K). One kit covers both hands."),
    ("Servo extension cables",
     "10pc 30cm 3-Pin Servo Extension Cable (JR, F-M)",
     "https://www.amazon.ca/dp/B09VGDJ6Z8", 8.27, 1,
     "Forearm servos to PCA9685 headers. Free shipping but slow (est. Oct 20-Nov 9); Prime alt in Alternatives."),
    ("Servo power wire",
     "BNTECHGO 18 AWG Silicone Wire, Red/Black, 10ft",
     "https://www.amazon.ca/dp/B01AQOI36M", 13.44, 1,
     "18AWG handles the 10A servo bus; Dupont wire (24AWG) must not carry servo power. 4.5* (438)."),
]

OPTIONAL = [
    ("Fingertip force sensors (x2, proof of concept)",
     "Interlink FSR 402 (30-81794), genuine - DigiKey.ca",
     "https://www.digikey.ca/en/products/detail/interlink-electronics/30-81794/2476468", 10.43, 2,
     "One per hand (e.g. index finger) on an ESP32 ADC pin via 10k divider. Qty-1 price break $10.43; "
     "qty 10 drops to $9.02/ea ($90.19) if you want all fingertips. DigiKey shipping for small orders NOT verified "
     "(free over $100 CAD). Amazon clones have 1 review - avoid."),
]

# (Item, Product Name, Link, Pack Price, Packs, Notes)
ALTERNATIVES = [
    ("Servos - BUDGET tier", "Miuzei MG90S Metal Gear Micro Servo, 10-Pack",
     "https://www.amazon.ca/dp/B0BWJ26PX2", 39.99, 1,
     "Saves $65.67 vs MG996R, but only ~2 kg-cm and micro size: needs a remixed (non-stock) InMoov forearm. 4.5* (784)."),
    ("Servos - cheapest metal-gear standard", "Aideepen MG996R, 6-Pack",
     "https://www.amazon.ca/dp/B09LS7RB5J", 39.99, 2,
     "12 servos for $79.98. Fallback if Deegoo sells out. Price from search results page. 4.3* (123)."),
    ("Servos - alt MG996R", "Miuzei MG996R Metal Gear, 4-Pack",
     "https://www.amazon.ca/dp/B0BZ4N367M", 42.99, 3,
     "Spec'd 15.2 kg-cm @6V. 4.4* (163)."),
    ("Servos - PREMIUM tier", "Miuzei 20KG DS3218 Metal Gear Digital, 180deg, 4-Pack",
     "https://www.amazon.ca/dp/B0FH53CKCS", 65.99, 3,
     "~20 kg-cm, stronger grip; $197.97 for 12 (+$92 vs core). Search results price. 4.6* (47)."),
    ("Servo driver - single board", "JZK PCA9685 16-Channel PWM Driver (single)",
     "https://www.amazon.ca/dp/B06XSFFXQY", 10.59, 1,
     "One board can run all 10 servos, but then one ESP32 drives both hands. 4.3* (222)."),
    ("Power - alt per-hand PSU", "SHNITPWR 3.5-12V 12A 144W Adjustable",
     "https://www.amazon.ca/dp/B0CKXW6ZCJ", 36.79, 2, "More headroom, same price. 4.1* (102)."),
    ("Power - ONE supply for both hands", "Aclorol 5V 20A 100W Switching PSU (screw terminals)",
     "https://www.amazon.ca/dp/B07KC55TJF", 39.44, 1,
     "Saves ~$34 but max ~5.75V (less torque), single point of failure, exposed mains terminals need a cover. "
     "Needs AC cord below. 4.5* (182)."),
    ("Power - AC cord for 20A PSU", "Bergen 3ft 16AWG 3-wire pigtail cord",
     "https://www.amazon.ca/dp/B07BQCMPF2", 6.87, 1, "Only with the single-PSU option. Search results price."),
    ("Barrel adapters - small pack", "2 Male + 2 Female 5.5x2.1 to Screw Terminal",
     "https://www.amazon.ca/dp/B07JMY5XXT", 9.35, 1, "Search results price. 4.6* (114)."),
    ("Tendon - alt", "Reaction Tackle Braided Line 50lb 150yd",
     "https://www.amazon.ca/dp/B01MRMI7R1", 21.41, 1, "4.3* (26.5K). Diameter not listed."),
    ("Springs - alt", "Glarks 100pc Extension Springs, 25 sizes x 4",
     "https://www.amazon.ca/dp/B0D7CP84F4", 20.95, 1, "Only 4 per size - not enough matching springs for 10 fingers."),
    ("Jumper kit - alt", "240pc 10cm+20cm Jumper Wire Kit",
     "https://www.amazon.ca/dp/B0BTT48V7P", 15.99, 1, "Search results price. 4.6* (551)."),
    ("Servo extensions - Prime alt", "10pc 30cm 3-Pin Servo Extension (Futaba/JR)",
     "https://www.amazon.ca/dp/B07NJB58XC", 11.99, 1, "Shipped by Amazon (faster). 4.4* (40)."),
    ("Power wire - longer", "BNTECHGO 18 AWG Silicone Wire 25ft",
     "https://www.amazon.ca/dp/B07HGTKQ89", 19.82, 1, "Search results price."),
    ("FSR - 10 genuine (all fingertips)", "Interlink FSR 402 (30-81794) x10 - DigiKey.ca",
     "https://www.digikey.ca/en/products/detail/interlink-electronics/30-81794/2476468", 90.19, 1,
     "$9.02/ea at qty 10; one per fingertip on both hands."),
    ("FSR - Amazon clone", "RELAND SUN FSR402-short (clone), sold singly",
     "https://www.amazon.ca/dp/B0B5F7JDT5", 8.19, 10, "Only 1 review, +$9.50 shipping. Not recommended."),
]

# Phase 2 full robot - USD, from research/physical_ai_research.md section 6
PHASE2 = [
    ("SO-ARM101 Pro motor kit (leader+follower servos, driver, cables)", "Seeed Studio",
     "https://www.seeedstudio.com/SO-101-Low-Cost-AI-Arm-Kit-Pro-p-6427.html", 249.90, 2,
     "Core arm. 2 kits = bimanual. STS3215 bus servos. LeRobot-native."),
    ("SO-101 printed parts", "Seeed Studio (or print yourself = $0)",
     "https://www.seeedstudio.com/SO-101-Low-Cost-AI-Arm-Kit-Pro-p-6427.html", 0, 2,
     "3D printing is covered, so $0. Seeed sells printed sets at $29.90 if needed."),
    ("Mobile base (LeKiwi base, or DIY XLeRobot IKEA cart)", "AIFITLAB LeKiwi",
     "https://aifitlab.com/products/lerobot-lekiwi-low-cost-mobile-manipulator", 260.00, 1,
     "'from $260'. XLeRobot full BOM ~$660 incl. arms: https://github.com/Vector-Wangel/XLeRobot"),
    ("USB cameras (head + 2 wrist)", "Seeed X10 USB camera",
     "https://www.seeedstudio.com/LeKiwi-Full-Kit-12V-Verision.html", 12.00, 3, "Listed as LeKiwi add-on."),
    ("On-robot compute", "NVIDIA Jetson Orin Nano Super Dev Kit",
     "https://jetsonhacks.com/2024/12/17/jetson-orin-nano-super-developer-kit/", 249.00, 1,
     "Runs ACT/SmolVLA locally. Skip and use a laptop to save $249."),
]


def build():
    wb = Workbook()
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title = Font(name="Arial", size=14, bold=True)
    hdr_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    sec_fill = PatternFill("solid", fgColor="D9E1F2")
    opt_fill = PatternFill("solid", fgColor="FFF2CC")
    tot_fill = PatternFill("solid", fgColor="E2EFDA")
    link_font = Font(name="Arial", size=10, color="0563C1", underline="single")
    input_font = Font(name="Arial", size=10, color="0000FF")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    money = '"$"#,##0.00'
    cols = ["Item", "Product Name", "Link", "Unit Price (CAD)", "Quantity", "Line Total (CAD)", "Notes/Rationale"]
    widths = [34, 46, 44, 15, 10, 16, 80]

    def header(ws, r, labels):
        for c, h in enumerate(labels, 1):
            cell = ws.cell(r, c, h)
            cell.font, cell.fill, cell.border = hdr_font, hdr_fill, box
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def row(ws, r, rec, fill=None):
        item, name, url, price, qty, note = rec
        vals = [item, name, url, price, qty, f"=D{r}*E{r}", note]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = base, box
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill
        ws.cell(r, 3).hyperlink = url
        ws.cell(r, 3).font = link_font
        ws.cell(r, 4).font = input_font
        ws.cell(r, 5).font = input_font
        ws.cell(r, 4).number_format = money
        ws.cell(r, 6).number_format = money

    def label_value(ws, r, label, formula, fill, fmt=money, font=bold):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        a = ws.cell(r, 1, label)
        a.font, a.fill = font, fill
        a.alignment = Alignment(horizontal="right")
        v = ws.cell(r, 6, formula)
        v.font, v.fill, v.number_format, v.border = bold, fill, fmt, box
        for c in range(1, 8):
            ws.cell(r, c).fill = fill

    def setup(ws, w=widths):
        for i, wd in enumerate(w, 1):
            ws.column_dimensions[chr(64 + i)].width = wd
        ws.sheet_view.showGridLines = False

    # ---------------- BOM sheet ----------------
    ws = wb.active
    ws.title = "BOM"
    setup(ws)
    ws["A1"] = "robotai - Phase 1: Two InMoov-style robotic hands - Bill of Materials"
    ws["A1"].font = title
    ws["A2"] = (f"All prices CAD, read from the linked pages on {VERIFIED}. 3D-printed parts excluded (printing covered). "
                "Blue cells = inputs (edit price/qty); totals recalc automatically.")
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    ws["A3"] = "Budget (CAD)"
    ws["A3"].font = bold
    ws["B3"] = 500
    ws["B3"].font, ws["B3"].number_format = Font(name="Arial", size=10, bold=True, color="0000FF"), money
    ws["C3"] = "Sales tax rate (ON HST)"
    ws["C3"].font = bold
    ws["D3"] = 0.13
    ws["D3"].font, ws["D3"].number_format = Font(name="Arial", size=10, color="0000FF"), "0%"

    r = 5
    ws.cell(r, 1, "CORE BOM").font = bold
    for c in range(1, 8):
        ws.cell(r, c).fill = sec_fill
    r += 1
    header(ws, r, cols)
    r += 1
    core_start = r
    for rec in CORE:
        row(ws, r, rec)
        r += 1
    core_end = r - 1
    sub = r
    label_value(ws, r, "Core subtotal (pre-tax)", f"=SUM(F{core_start}:F{core_end})", tot_fill); r += 1
    tax = r
    label_value(ws, r, "Est. HST on core (13%, shipping not included)", f"=F{sub}*$D$3", tot_fill, font=base); r += 1
    tot = r
    label_value(ws, r, "CORE TOTAL incl. tax", f"=F{sub}+F{tax}", tot_fill); r += 1
    label_value(ws, r, "Headroom under budget", f"=$B$3-F{tot}", tot_fill); r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(r, 1, "Under $500 CAD?").font = bold
    ws.cell(r, 1).alignment = Alignment(horizontal="right")
    v = ws.cell(r, 6, f'=IF(F{tot}<=$B$3,"YES - under budget","NO - over budget")')
    v.font = Font(name="Arial", size=10, bold=True, color="006100")
    for c in range(1, 8):
        ws.cell(r, c).fill = tot_fill
    r += 2

    ws.cell(r, 1, "OPTIONAL - fingertip force sensors (NOT in core total)").font = bold
    for c in range(1, 8):
        ws.cell(r, c).fill = opt_fill
    r += 1
    header(ws, r, cols)
    r += 1
    opt_start = r
    for rec in OPTIONAL:
        row(ws, r, rec, fill=opt_fill)
        r += 1
    opt_end = r - 1
    osub = r
    label_value(ws, r, "Optional subtotal (pre-tax)", f"=SUM(F{opt_start}:F{opt_end})", opt_fill); r += 1
    label_value(ws, r, "Core + optional, incl. tax", f"=(F{sub}+F{osub})*(1+$D$3)", opt_fill); r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(r, 1, "Still under $500 with sensors?").font = bold
    ws.cell(r, 1).alignment = Alignment(horizontal="right")
    ws.cell(r, 6, f'=IF(F{r-1}<=$B$3,"YES","NO")').font = bold
    for c in range(1, 8):
        ws.cell(r, c).fill = opt_fill
    r += 2

    notes = [
        "RECOMMENDATION: buy the core list + the 2 FSRs. Both fit comfortably under $500 CAD with tax, leaving room for shipping and a spare servo pack.",
        "If you later want all 10 fingertips sensed, the DigiKey 10-pack ($90.19) still keeps core + sensors under $500 with tax.",
        "Wiring rule: ESP32 3.3V/GND -> PCA9685 VCC/GND, SDA/SCL -> GPIO21/22. PSU (6V) -> PCA9685 V+ screw terminal. Common ground between PSU and ESP32. Never feed servos from the ESP32 5V pin.",
        "Not included: shipping (most Amazon.ca items ship free with Prime or over $35), tools (soldering iron, crimper), and fasteners/bearings from the InMoov print list.",
        "Prices change daily - re-check the links before ordering. See 'Alternatives' for budget/premium swaps and 'Phase 2 - Full robot' for arms/base.",
    ]
    ws.cell(r, 1, "Notes").font = bold
    r += 1
    for n in notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        c = ws.cell(r, 1, n)
        c.font = base
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 28
        r += 1
    ws.freeze_panes = "A7"

    # ---------------- Alternatives ----------------
    wa = wb.create_sheet("Alternatives")
    setup(wa)
    wa["A1"] = "Alternatives & swaps (not in any total unless you swap them in)"
    wa["A1"].font = title
    alt_cols = ["Item", "Product Name", "Link", "Pack Price (CAD)", "Packs", "Line Total (CAD)", "Notes/Rationale"]
    header(wa, 3, alt_cols)
    for i, rec in enumerate(ALTERNATIVES, 4):
        row(wa, i, rec)

    # ---------------- Phase 2 ----------------
    wp = wb.create_sheet("Phase 2 - Full robot")
    setup(wp)
    wp["A1"] = "Phase 2 - bimanual mobile robot (XLeRobot-style), PRICES IN USD"
    wp["A1"].font = title
    wp["A2"] = ("From research/physical_ai_research.md s.6. Separate from the $500 CAD hands budget. "
                "Prices from vendor pages/coverage on 2026-09-25.")
    wp["A2"].font = Font(name="Arial", size=9, italic=True)
    wp["C3"] = "USD->CAD rate (assumed, verify)"
    wp["C3"].font = bold
    wp["D3"] = 1.38
    wp["D3"].font = input_font
    p_cols = ["Item", "Product / Vendor", "Link", "Unit Price (USD)", "Quantity", "Line Total (USD)", "Notes/Rationale"]
    header(wp, 5, p_cols)
    r = 6
    for rec in PHASE2:
        row(wp, r, rec)
        wp.cell(r, 4).number_format = '"US$"#,##0.00'
        wp.cell(r, 6).number_format = '"US$"#,##0.00'
        r += 1
    label_value(wp, r, "Phase 2 subtotal (USD, pre-tax/shipping)", f"=SUM(F6:F{r-1})", tot_fill,
                fmt='"US$"#,##0.00'); r += 1
    label_value(wp, r, "Phase 2 subtotal (CAD, approx.)", f"=F{r-1}*$D$3", tot_fill); r += 2
    wp.cell(r, 1, ("Cerebras = System-2 brain (planner, voice, replanning, success checks) in the cloud; "
                   "ACT/SmolVLA motor policy runs locally at 30-50 Hz. See research brief.")).font = base

    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
